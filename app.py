import re
from io import BytesIO

import streamlit as st
from docx import Document
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

st.set_page_config(page_title="Resume & Job Matching Assistant", page_icon="📄", layout="wide")

st.title("📄 Resume & Job Matching Assistant")
st.write(
    "Upload a resume and paste a job description to estimate how closely they match. "
    "This is a helpful screening aid, not a hiring decision or a guarantee of job fit."
)

STOP_WORDS = {
    "a", "an", "and", "are", "as", "at", "be", "by", "for", "from", "in", "into",
    "is", "it", "of", "on", "or", "our", "that", "the", "this", "to", "with",
    "you", "your", "will", "we", "they", "their", "have", "has", "can", "must",
    "required", "preferred", "ability", "work", "working", "years", "year"
}

def extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)

def extract_docx(file_bytes: bytes) -> str:
    doc = Document(BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)

def extract_resume(uploaded_file) -> str:
    data = uploaded_file.getvalue()
    suffix = uploaded_file.name.lower()
    if suffix.endswith(".pdf"):
        return extract_pdf(data)
    if suffix.endswith(".docx"):
        return extract_docx(data)
    raise ValueError("Please upload a PDF or DOCX resume.")

def keywords(text: str) -> set[str]:
    words = re.findall(r"[a-zA-Z][a-zA-Z+#.]{1,}", text.lower())
    return {w.strip(".") for w in words if len(w.strip(".")) > 2 and w not in STOP_WORDS}

def calculate_match(resume_text: str, job_text: str):
    vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2))
    matrix = vectorizer.fit_transform([resume_text, job_text])
    score = float(cosine_similarity(matrix[0:1], matrix[1:2])[0][0] * 100)

    resume_terms = keywords(resume_text)
    job_terms = keywords(job_text)
    matched = sorted(resume_terms & job_terms)
    missing = sorted(job_terms - resume_terms)
    return score, matched, missing

with st.sidebar:
    st.header("How to use")
    st.markdown(
        "1. Upload a PDF or DOCX resume.\n"
        "2. Paste the job description.\n"
        "3. Click **Analyze match**.\n"
        "4. Review the score and keyword gaps."
    )
    st.caption("Privacy tip: use a resume you have permission to process. This demo analyzes text in the app session and does not save it to a database.")

left, right = st.columns([1, 1])
with left:
    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx"])
with right:
    job_text = st.text_area(
        "Paste the job description",
        height=220,
        placeholder="Paste the responsibilities, skills, and qualifications from the job post here..."
    )

if st.button("Analyze match", type="primary", use_container_width=True):
    if not uploaded:
        st.error("Please upload a resume first.")
    elif not job_text.strip():
        st.error("Please paste a job description first.")
    else:
        try:
            resume_text = extract_resume(uploaded)
            if len(resume_text.strip()) < 30:
                st.error("Very little text was extracted. If this is a scanned PDF, try a text-based PDF or DOCX file.")
            else:
                score, matched, missing = calculate_match(resume_text, job_text)
                st.subheader("Match overview")
                metric_col, matched_col = st.columns([1, 2])
                with metric_col:
                    st.metric("Text similarity", f"{score:.1f}%")
                    st.progress(min(max(int(score), 0), 100))
                with matched_col:
                    st.write("This score measures text similarity using TF-IDF. It is not a probability of getting hired.")
                    st.write(f"**Potentially matching keywords:** {len(matched)}")
                    st.write(f"**Job-description keywords not found in resume:** {len(missing)}")

                st.divider()
                st.subheader("Keywords found in both")
                if matched:
                    st.write(", ".join(matched[:80]))
                else:
                    st.info("No shared keywords were detected. Try checking that the resume text extracted correctly.")

                st.subheader("Potential keywords to review")
                st.caption("Only add skills or terms that truthfully reflect your experience.")
                if missing:
                    st.write(", ".join(missing[:80]))
                else:
                    st.success("The simple keyword scan did not find any gaps.")

                st.subheader("Resume improvement checklist")
                checklist = [
                    "Tailor the summary/objective to the role.",
                    "Move the most relevant skills and experience closer to the top.",
                    "Use measurable outcomes where accurate (for example, time saved or projects delivered).",
                    "Include relevant tools and technologies only if you can discuss them honestly.",
                    "Proofread dates, spelling, formatting, and contact details."
                ]
                for item in checklist:
                    st.checkbox(item, key=f"check_{item}")

                report = (
                    "RESUME & JOB MATCHING REPORT\n"
                    f"Text similarity: {score:.1f}%\n\n"
                    f"Keywords found in both ({len(matched)}):\n"
                    + ", ".join(matched)
                    + "\n\nPotential keywords to review ("
                    + str(len(missing))
                    + "):\n"
                    + ", ".join(missing)
                    + "\n\nNote: This report is based on text similarity and simple keyword overlap. "
                    "It is not a hiring recommendation or a probability of employment.\n"
                )
                st.download_button(
                    "Download analysis report (.txt)",
                    data=report.encode("utf-8"),
                    file_name="resume_job_match_report.txt",
                    mime="text/plain",
                    use_container_width=True
                )
        except Exception as exc:
            st.error(f"Could not analyze this file: {exc}")
            st.caption("Try a text-based PDF or DOCX file.")

st.divider()
st.caption("Built with Python, Streamlit, scikit-learn, pypdf, and python-docx.")
